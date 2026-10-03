"""
Tadqiqot eksporti: anonim yozuvlar (CSV/Excel) + tayyor jadvallar (Excel).
Ismli ro'yxat — alohida, faqat admin uchun (build_patients_xlsx).

SPSS / R / Python'ga to'g'ridan-to'g'ri import qilinadi. Ism va Telegram ID
chiqmaydi — faqat subject_id (S0001…). Kodlash "Lug'at" varag'ida.
"""

import csv
from io import BytesIO, StringIO

from app.calculator import HOMA_IR_CUTOFF, classify_homa_ir
from app.stats import fmt_desc, fmt_p, subject_key
from app.utils import TASHKENT, parse_utc

COLUMNS: list[tuple[str, str]] = [
    ("subject_id", "Sub'ekt kodi (anonim)"),
    ("visit_no", "Shu sub'ektning nechanchi o'lchovi (1 = baseline)"),
    ("date", "Sana (Toshkent vaqti)"),
    ("mode", "full — to'liq skrining, quick — tezkor HOMA-IR"),
    ("sex", "M — erkak, F — ayol"),
    ("age", "Yosh, yil"),
    ("weight_kg", "Vazn, kg"),
    ("height_cm", "Bo'y, sm"),
    ("waist_cm", "Bel aylanasi, sm"),
    ("bmi", "Tana massasi indeksi, kg/m²"),
    ("whtr", "Bel/bo'y nisbati"),
    ("activity_daily", "Kuniga ≥30 daqiqa faollik: 1 ha, 0 yo'q"),
    ("vegetables_daily", "Har kuni sabzavot/meva: 1 ha, 0 yo'q"),
    ("bp_medication", "Antigipertenziv dori: 1 ha, 0 yo'q"),
    ("history_high_glucose", "Ilgari yuqori glukoza: 1 ha, 0 yo'q"),
    ("family_diabetes", "Oilaviy anamnez: none / second / first"),
    ("glucose_mmol", "Och qoringa glukoza, mmol/L"),
    ("insulin_uiu", "Och qoringa insulin, μU/mL"),
    ("homa_ir", "HOMA-IR"),
    ("homa_beta", "HOMA-β, %"),
    ("quicki", "QUICKI"),
    ("homa_class", "HOMA-IR toifasi"),
    ("ir", f"Insulin rezistentligi (HOMA-IR ≥ {HOMA_IR_CUTOFF}): 1/0"),
    ("findrisc", "FINDRISC ball (0–26)"),
    ("findrisc_band", "FINDRISC toifasi"),
]


def anonymized_records(rows: list[dict], anon_only: bool = False) -> list[dict]:
    """Xronologik tartibda; subject_id birinchi paydo bo'lish tartibida beriladi."""
    ids: dict[str, str] = {}
    visits: dict[str, int] = {}
    out = []
    for r in sorted(rows, key=lambda x: (str(x.get("created_at")), x.get("id", 0))):
        if anon_only and not r.get("is_anonymous"):
            continue
        key = subject_key(r)
        if key not in ids:
            ids[key] = f"S{len(ids) + 1:04d}"
        visits[key] = visits.get(key, 0) + 1
        homa = r.get("homa_ir")
        out.append({
            "subject_id": ids[key],
            "visit_no": visits[key],
            "date": parse_utc(r["created_at"]).astimezone(TASHKENT).strftime("%Y-%m-%d %H:%M"),
            "mode": r.get("kind") or "full",
            "sex": r.get("sex"),
            "age": r.get("age"),
            "weight_kg": r.get("weight"),
            "height_cm": r.get("height"),
            "waist_cm": r.get("waist"),
            "bmi": round(r["bmi"], 2) if r.get("bmi") is not None else None,
            "whtr": round(r["whtr"], 3) if r.get("whtr") is not None else None,
            "activity_daily": r.get("activity_yes"),
            "vegetables_daily": r.get("veg_daily"),
            "bp_medication": r.get("bp_meds"),
            "history_high_glucose": r.get("high_glucose_hist"),
            "family_diabetes": r.get("family_hx"),
            "glucose_mmol": r.get("fasting_glucose"),
            "insulin_uiu": r.get("fasting_insulin"),
            "homa_ir": round(homa, 3) if homa is not None else None,
            "homa_beta": round(r["homa_beta"], 1) if r.get("homa_beta") is not None else None,
            "quicki": round(r["quicki"], 4) if r.get("quicki") is not None else None,
            "homa_class": classify_homa_ir(homa) if homa is not None else None,
            "ir": (1 if homa >= HOMA_IR_CUTOFF else 0) if homa is not None else None,
            "findrisc": r.get("findrisc"),
            "findrisc_band": r.get("findrisc_band"),
        })
    return out


def build_csv(rows: list[dict], anon_only: bool = False) -> bytes:
    buf = StringIO()
    w = csv.writer(buf)
    w.writerow([c for c, _ in COLUMNS])
    for rec in anonymized_records(rows, anon_only):
        w.writerow(["" if rec[c] is None else rec[c] for c, _ in COLUMNS])
    # BOM — Excel o'zbek/kirill harflarini to'g'ri ochishi uchun
    return ("﻿" + buf.getvalue()).encode("utf-8")


def build_xlsx(rows: list[dict], rep: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="1E3A8A")

    def sheet(ws, header: list[str], data: list[list], widths: list[int] | None = None):
        ws.append(header)
        for cell in ws[1]:
            cell.font, cell.fill = head_font, head_fill
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        for row in data:
            ws.append(row)
        ws.freeze_panes = "A2"
        for i, w in enumerate(widths or [], start=1):
            ws.column_dimensions[get_column_letter(i)].width = w

    wb = Workbook()
    ws = wb.active
    ws.title = "Ma'lumotlar"
    recs = anonymized_records(rows)
    sheet(ws, [c for c, _ in COLUMNS], [[r[c] for c, _ in COLUMNS] for r in recs],
          [11, 8, 17, 7] + [11] * (len(COLUMNS) - 4))
    ws.auto_filter.ref = ws.dimensions

    t1 = []
    for t in rep["table1"]:
        row = [t["name"] + (f", {t['unit']}" if t["unit"] else "")]
        for grp in ("all", "M", "F"):
            d = t[grp]
            msd, med = fmt_desc(d, t["digits"]) if d.get("n") else ("—", "—")
            row += [d.get("n", 0), msd, med]
        row.append(fmt_p(t["p"]))
        t1.append(row)
    sheet(wb.create_sheet("1-jadval"),
          ["Ko'rsatkich", "Umumiy n", "Umumiy M±SD", "Umumiy Me [Q1; Q3]",
           "Erkak n", "Erkak M±SD", "Erkak Me [Q1; Q3]",
           "Ayol n", "Ayol M±SD", "Ayol Me [Q1; Q3]", "p (Mann–Whitney)"],
          t1, [26, 9, 16, 22, 9, 16, 22, 9, 16, 22, 14])

    prev = []
    if rep["n_labs"]:
        o = rep["ir_overall"]
        prev.append(["Umumiy", "", o["k"], o["n"], round(o["pct"], 1), round(o["lo"], 1), round(o["hi"], 1), ""])
        for title, block in (("Jins", rep["ir_by_sex"]), ("Yosh guruhi", rep["ir_by_age"]),
                             ("BMI toifasi", rep["ir_by_bmi"])):
            for j, it in enumerate(i for i in block["items"] if i["n"]):
                prev.append([title, it["label"], it["k"], it["n"], round(it["pct"], 1),
                             round(it["lo"], 1), round(it["hi"], 1), fmt_p(block["p"]) if j == 0 else ""])
    sheet(wb.create_sheet("IR prevalentligi"),
          ["Kesim", "Guruh", "IR (n)", "Jami (n)", "%", "95% CI past", "95% CI yuqori", "p"],
          prev, [14, 22, 9, 9, 8, 12, 13, 9])

    sheet(wb.create_sheet("Korrelyatsiya"),
          ["HOMA-IR bilan", "n", "Spearman ρ", "p"],
          [[c["name"], c["n"], None if c["rho"] is None else round(c["rho"], 3), fmt_p(c["p"])]
           for c in rep["correlations"]], [26, 8, 12, 10])

    ref = rep["reference"]
    d = rep.get("homa_desc", {})
    cut = [
        ["Butun namuna n", d.get("n", 0)],
        ["Q1", round(d["q1"], 3) if d.get("n") else None],
        ["Mediana", round(d["median"], 3) if d.get("n") else None],
        ["Q3", round(d["q3"], 3) if d.get("n") else None],
        ["Sog'lom guruh n", ref["n"]],
        ["Sog'lom guruh P75 (cutoff)", None if ref.get("p75") is None else round(ref["p75"], 3)],
        ["Sog'lom guruh P90", None if ref.get("p90") is None else round(ref["p90"], 3)],
        ["Ishonchli (n ≥ 30)", "ha" if ref.get("reliable") else "yo'q"],
    ]
    sheet(wb.create_sheet("Cutoff"), ["Ko'rsatkich", "Qiymat"], cut, [30, 14])

    lo = rep["longitudinal"]
    dyn = [["Sub'ektlar (≥2 o'lchov)", lo.get("n", 0)]]
    if lo.get("n"):
        dyn += [
            ["Birinchi HOMA-IR, Me", round(lo["first"]["median"], 3)],
            ["Oxirgi HOMA-IR, Me", round(lo["last"]["median"], 3)],
            ["Δ HOMA-IR, Me", round(lo["delta"]["median"], 3)],
            ["≥10% yaxshilangan", lo["improved"]],
            ["≥10% yomonlashgan", lo["worsened"]],
            ["Wilcoxon p", fmt_p(lo["p"])],
        ]
    sheet(wb.create_sheet("Dinamika"), ["Ko'rsatkich", "Qiymat"], dyn, [28, 14])

    sheet(wb.create_sheet("Lug'at"), ["Ustun", "Ma'nosi"], [list(c) for c in COLUMNS], [22, 70])

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


PATIENT_COLUMNS: list[tuple[str, str, int]] = [
    # (kalit, sarlavha, ustun kengligi)
    ("subject_id", "Kod (anonim faylda)", 12),
    ("full_name", "Ism-familiya", 28),
    ("telegram", "Telegram", 18),
    ("age", "Yosh", 7),
    ("sex", "Jins", 7),
    ("date", "Sana", 17),
    ("visit_no", "O'lchov №", 10),
    ("mode", "Rejim", 9),
    ("glucose_mmol", "Glukoza, mmol/L", 14),
    ("insulin_uiu", "Insulin, μU/mL", 14),
    ("homa_ir", "HOMA-IR", 10),
    ("homa_class", "HOMA-IR toifasi", 26),
    ("bmi", "BMI", 8),
    ("waist_cm", "Bel, sm", 8),
    ("findrisc", "FINDRISC", 10),
    ("findrisc_band", "FINDRISC toifasi", 16),
]


def build_patients_xlsx(rows: list[dict]) -> bytes:
    """
    FAQAT admin uchun: ism-familiya va yosh bilan bemorlar ro'yxati.
    subject_id anonim fayldagi kod bilan bir xil — ikkalasini bog'lash mumkin.
    Bu fayl dissertatsiya/jurnalga yuborilmaydi (shaxsiy ma'lumot).
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    ordered = sorted(rows, key=lambda x: (str(x.get("created_at")), x.get("id", 0)))
    recs = anonymized_records(ordered)
    wb = Workbook()
    ws = wb.active
    ws.title = "Bemorlar"
    ws.append([title for _, title, _ in PATIENT_COLUMNS])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="991B1B")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for raw, rec in zip(ordered, recs):
        rec = dict(rec, full_name=raw.get("patient_name") or "",
                   telegram=f"@{raw['tg_username']}" if raw.get("tg_username") else "")
        ws.append([rec.get(key) for key, _, _ in PATIENT_COLUMNS])
    for i, (_, _, width) in enumerate(PATIENT_COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions
    note = wb.create_sheet("Eslatma")
    note["A1"] = "⚠️ Shaxsiy tibbiy ma'lumot. Faqat tadqiqotchi uchun."
    note["A2"] = "Dissertatsiya, jurnal va boshqalarga anonim Excel/CSV faylni yuboring."
    note["A3"] = "\"Kod\" ustuni anonim fayldagi subject_id bilan bir xil."
    note["A1"].font = Font(bold=True, color="991B1B")
    note.column_dimensions["A"].width = 80
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
