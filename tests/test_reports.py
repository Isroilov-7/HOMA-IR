from io import BytesIO

from app import charts, stats
from app.calculator import calculate_quicki, quick_report
from app.reports.patient_pdf import build_pdf_report
from app.reports.research_pdf import build_research_pdf
from app.reports.research_xlsx import anonymized_records, build_csv, build_xlsx
from conftest import make_rows


def _is_pdf(b: bytes) -> bool:
    return b[:5] == b"%PDF-" and len(b) > 2000


def _is_png(b: bytes) -> bool:
    return b[:8] == b"\x89PNG\r\n\x1a\n"


def test_patient_pdf_full_quick_and_cyrillic():
    rows = make_rows(3)
    full = dict(rows[0], patient_name="Алиев Ғани Oʻgʻli")
    hist = [dict(rows[1], id=2, patient_name=full["patient_name"]), full]
    png = charts.patient_trend_png(hist)
    assert png and _is_png(png)
    assert _is_pdf(build_pdf_report({"display_name": "x"}, full, hist, png))

    rep = quick_report(glucose=5.0, insulin=12)
    quick = {"id": 9, "kind": "quick", "patient_name": "Vali", "age": 40, "fasting_glucose": 5.0,
             "fasting_insulin": 12.0, "homa_ir": rep["homa_ir"], "quicki": calculate_quicki(5.0, 12),
             "created_at": "2026-10-02 05:00:00"}
    assert _is_pdf(build_pdf_report(None, quick))


def test_research_pdf_and_charts():
    rows = make_rows(60)
    rep = stats.cohort_report(rows)
    pdf = build_research_pdf(rows, rep)
    assert _is_pdf(pdf) and len(pdf) > 50_000  # rasmlar bilan
    assert _is_png(charts.findrisc_bars_png(rep["findrisc_bands"]))


def test_research_pdf_empty_db_does_not_crash():
    assert _is_pdf(build_research_pdf([], stats.cohort_report([])))


def test_anonymized_export_has_no_names_or_ids():
    rows = make_rows(10)
    recs = anonymized_records(rows)
    assert recs[0]["subject_id"] == "S0001"
    blob = build_csv(rows).decode("utf-8-sig")
    assert "Bemor" not in blob and "1000" not in blob.split("\n", 1)[1].split(",")[0]
    visits = [r for r in recs if r["visit_no"] == 2]
    assert len(visits) == 10  # make_rows: birinchi 15 ta (bu yerda 10 ta) — ikki marta


def test_xlsx_has_all_sheets():
    from openpyxl import load_workbook

    rows = make_rows(40)
    wb = load_workbook(BytesIO(build_xlsx(rows, stats.cohort_report(rows))))
    assert wb.sheetnames == ["Ma'lumotlar", "1-jadval", "IR prevalentligi", "Korrelyatsiya",
                             "Cutoff", "Dinamika", "Lug'at"]
    assert wb["Ma'lumotlar"].max_row == len(rows) + 1


def test_research_pdf_small_sample_unreliable_cutoff():
    rows = make_rows(40)  # sog'lom guruh n=8: cutoff hisoblanadi, lekin "taxminiy"
    rep = stats.cohort_report(rows)
    assert rep["reference"]["p75"] is not None and not rep["reference"]["reliable"]
    assert _is_pdf(build_research_pdf(rows, rep))
